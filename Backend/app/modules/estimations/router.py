import os
from fastapi import APIRouter, Depends, UploadFile, File, Form, HTTPException, status
from fastapi.responses import FileResponse, RedirectResponse
from sse_starlette.sse import EventSourceResponse
from app.dependencies import get_current_user, get_db
from typing import Union
from app.modules.estimations.schemas import DraftSessionRequest, NotificationRequest, EncryptedPayloadSchema, CompleteSessionRequest
from app.core.crypto_utils import decrypt_aes_key_with_rsa, decrypt_payload_with_aes_gcm
from app.modules.estimations.service import estimation_service
from app.modules.estimations.repository import estimation_repo
from app.services.graph.session_manager import session_manager
from app.services.graph.graph import costmate_graph
from app.config import settings

router = APIRouter(prefix="/api", tags=["Estimations"])

@router.post("/upload")
async def upload_floor_plan(file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
    return await estimation_service.upload_floor_plan(current_user["id"], file)

@router.get("/session")
async def list_user_sessions(conn = Depends(get_db), current_user: dict = Depends(get_current_user)):
    return estimation_service.get_user_sessions(conn, current_user["id"])

@router.get("/session/{session_id}")
async def get_session_state(session_id: str, current_user: dict = Depends(get_current_user)):
    return await estimation_service.get_session_state(session_id, current_user["id"])

@router.post("/session/draft")
async def create_draft_session(req: DraftSessionRequest, conn = Depends(get_db), current_user: dict = Depends(get_current_user)):
    return estimation_service.create_draft_session(conn, current_user["id"], req)

@router.get("/session/draft/{session_id}")
async def get_draft_session(session_id: str, conn = Depends(get_db), current_user: dict = Depends(get_current_user)):
    return estimation_service.get_draft_session(conn, session_id, current_user["id"])

@router.post("/session/draft/{session_id}/upload")
async def upload_draft_file(session_id: str, file: UploadFile = File(...), conn = Depends(get_db), current_user: dict = Depends(get_current_user)):
    return await estimation_service.upload_draft_file(conn, session_id, current_user["id"], file)

@router.post("/session/draft/{session_id}/schedule")
async def upload_draft_schedule(session_id: str, file: UploadFile = File(...), conn = Depends(get_db), current_user: dict = Depends(get_current_user)):
    return await estimation_service.upload_draft_schedule(conn, session_id, current_user["id"], file)

@router.post("/session/draft/{session_id}/crop-schedule")
async def crop_schedule(
    session_id: str,
    file: UploadFile = File(...),
    x0: float = Form(...),
    y0: float = Form(...),
    x1: float = Form(...),
    y1: float = Form(...),
    page_num: int = Form(...),
    conn = Depends(get_db),
    current_user: dict = Depends(get_current_user)
):
    return await estimation_service.crop_schedule(conn, session_id, current_user["id"], file, x0, y0, x1, y1, page_num)

@router.post("/session/draft/{session_id}/specification")
async def upload_draft_specification(session_id: str, file: UploadFile = File(...), conn = Depends(get_db), current_user: dict = Depends(get_current_user)):
    return await estimation_service.upload_draft_specification(conn, session_id, current_user["id"], file)

@router.post("/session/draft/{session_id}/quick-scan")
async def quick_scan_draft(session_id: str, conn = Depends(get_db), current_user: dict = Depends(get_current_user)):
    return await estimation_service.quick_scan_draft(conn, session_id, current_user["id"])

from fastapi import Request

@router.post("/session/draft/{session_id}/complete")
async def complete_draft_session(session_id: str, request: Request, conn = Depends(get_db), current_user: dict = Depends(get_current_user)):
    try:
        body_bytes = await request.body()
        print(f"DEBUG COMPLETE ENDPOINT: {body_bytes[:200]}")
        import json
        req_data = {}
        if body_bytes:
            req_data = json.loads(body_bytes)
            # Check if encrypted manually in case middleware failed
            if "rsa_encrypted_aes_key" in req_data and "aes_encrypted_payload" in req_data:
                aes_key = decrypt_aes_key_with_rsa(req_data["rsa_encrypted_aes_key"])
                req_data = decrypt_payload_with_aes_gcm(req_data["aes_encrypted_payload"], aes_key)
        
        intake = req_data.get("intake_data") if isinstance(req_data, dict) else None
        return estimation_service.complete_draft_session(conn, session_id, current_user["id"], intake)
    except Exception as e:
        print(f"DEBUG COMPLETE ENDPOINT ERROR: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=400, detail=str(e))

@router.delete("/session/{session_id}")
async def delete_session(session_id: str, conn = Depends(get_db), current_user: dict = Depends(get_current_user)):
    session_manager.cancel_session(session_id)
    estimation_repo.delete_session(conn, session_id, current_user["id"])
    conn.commit()
    return {"status": "success"}

@router.post("/qa/{session_id}/submit")
async def submit_qa(session_id: str, payload: Union[EncryptedPayloadSchema, dict], current_user: dict = Depends(get_current_user)):
    if isinstance(payload, EncryptedPayloadSchema):
        aes_key = decrypt_aes_key_with_rsa(payload.rsa_encrypted_aes_key)
        decrypted_json = decrypt_payload_with_aes_gcm(payload.aes_encrypted_payload, aes_key)
        verified_qa = decrypted_json
    else:
        verified_qa = payload
    return await estimation_service.submit_qa(session_id, verified_qa, current_user["id"])

@router.get("/status/{session_id}/stream")
async def stream_session_status(session_id: str, current_user: dict = Depends(get_current_user)):
    restored = await session_manager.restore_session_if_needed(session_id, user_id=current_user["id"])
    if not restored:
        raise HTTPException(status_code=404, detail="Session not found.")
    return EventSourceResponse(session_manager.get_status_stream(session_id))

@router.get("/download/{session_id}")
async def download_output(session_id: str, current_user: dict = Depends(get_current_user)):
    restored = await session_manager.restore_session_if_needed(session_id, user_id=current_user["id"])
    if not restored:
        raise HTTPException(status_code=404, detail="Session not found.")
    config = {"configurable": {"thread_id": session_id}}
    state_snapshot = await costmate_graph.aget_state(config)
    if not state_snapshot.values:
        raise HTTPException(status_code=404, detail="Session not found.")
    
    # Regenerate the Excel file dynamically to use the latest code and state
    from app.services.agents.layer5_output.excel_writer_node import excel_writer_node
    from app.core.logging import logger
    try:
        excel_result = await excel_writer_node(state_snapshot.values)
        excel_path = excel_result.get("excel_file_path")
    except Exception as e:
        logger.error(f"Failed to regenerate Excel on download: {e}")
        excel_path = state_snapshot.values.get("excel_file_path")
        
    if not excel_path:
        raise HTTPException(status_code=404, detail="Excel output file not ready or not found.")
        
    filename = f"Costmate_Estimate_{session_id}.xlsx"
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    
    if excel_path.startswith("http"):
        import httpx
        from fastapi.responses import StreamingResponse
        # Proxy the Cloudinary URL to bypass frontend CORS issues for fetch()
        async def stream_external_file():
            async with httpx.AsyncClient(follow_redirects=True) as client:
                async with client.stream("GET", excel_path) as response:
                    if response.status_code != 200:
                        raise HTTPException(status_code=404, detail="Failed to fetch Excel from cloud storage")
                    async for chunk in response.aiter_bytes():
                        yield chunk
        
        return StreamingResponse(
            stream_external_file(),
            media_type=media_type,
            headers={
                "Content-Disposition": f"attachment; filename={filename}",
                "Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"
            }
        )
        
    if not os.path.exists(excel_path):
        raise HTTPException(status_code=404, detail="Excel output file not ready or not found.")
        
    return FileResponse(
        path=excel_path, 
        media_type=media_type, 
        filename=filename,
        headers={"Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"}
    )

@router.post("/download/{session_id}/custom-excel")
async def download_custom_excel(
    session_id: str,
    payload: dict,
    current_user: dict = Depends(get_current_user)
):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from fastapi.responses import FileResponse
    import tempfile
    
    sheets_data = payload.get("sheets", [])
    if not sheets_data:
        raise HTTPException(status_code=400, detail="No sheet data provided.")
        
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    for s in sheets_data:
        sheet_name = s.get("name", "Sheet1")
        ws = wb.create_sheet(title=sheet_name)
        
        celldata = s.get("celldata", [])
        data_matrix = s.get("data", [])
        
        # Prioritize live 2D grid matrix `data_matrix` which contains user edits from FortuneSheet
        if data_matrix and len(data_matrix) > 0:
            for r_idx, row in enumerate(data_matrix, 1):
                if not row: continue
                for c_idx, cell_obj in enumerate(row, 1):
                    if cell_obj is None:
                        continue
                    val = ""
                    font_bold = False
                    font_color = "000000"
                    bg_color = None
                    
                    if isinstance(cell_obj, dict):
                        val = cell_obj.get("m") if cell_obj.get("m") is not None else cell_obj.get("v", "")
                        if cell_obj.get("bl"): font_bold = True
                        if cell_obj.get("fc"): font_color = str(cell_obj.get("fc")).replace("#", "")
                        if cell_obj.get("bg"): bg_color = str(cell_obj.get("bg")).replace("#", "")
                    elif cell_obj is not None:
                        val = str(cell_obj)
                        
                    val_str = str(val) if val is not None else ""
                    if val_str != "" or (bg_color and len(bg_color) == 6) or r_idx == 1:
                        cell = ws.cell(row=r_idx, column=c_idx, value=val_str)
                        fill_color = "FFFFFF"
                        if bg_color and len(bg_color) == 6:
                            fill_color = bg_color
                        elif r_idx == 1:
                            fill_color = "333333"
                            
                        fc = "FFFFFF" if r_idx == 1 else (font_color if len(font_color) == 6 else "000000")
                        cell.font = Font(bold=(font_bold or r_idx == 1), color=fc)
                        if r_idx == 1:
                            cell.alignment = Alignment(horizontal="center")
                            cell.fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
                        elif bg_color and len(bg_color) == 6:
                            cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        elif celldata:
            for cell_item in celldata:
                r = cell_item.get("r", 0) + 1
                c = cell_item.get("c", 0) + 1
                v_obj = cell_item.get("v")
                
                val = ""
                font_bold = False
                font_color = "000000"
                bg_color = None
                
                if isinstance(v_obj, dict):
                    val = v_obj.get("m") if v_obj.get("m") is not None else v_obj.get("v", "")
                    if v_obj.get("bl"):
                        font_bold = True
                    if v_obj.get("fc"):
                        font_color = str(v_obj.get("fc")).replace("#", "")
                    if v_obj.get("bg"):
                        bg_color = str(v_obj.get("bg")).replace("#", "")
                elif v_obj is not None:
                    val = str(v_obj)
                    
                cell = ws.cell(row=r, column=c, value=str(val) if val is not None else "")
                
                fill_color = "FFFFFF"
                if bg_color and len(bg_color) == 6:
                    fill_color = bg_color
                elif r == 1:
                    fill_color = "333333"
                    
                fc = "FFFFFF" if r == 1 else (font_color if len(font_color) == 6 else "000000")
                
                cell.font = Font(bold=(font_bold or r == 1), color=fc)
                if r == 1:
                    cell.alignment = Alignment(horizontal="center")
                    cell.fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
                elif bg_color and len(bg_color) == 6:
                    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
                    
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    temp_dir = tempfile.gettempdir()
    output_filename = f"Costmate_Estimate_{session_id}.xlsx"
    temp_path = os.path.join(temp_dir, output_filename)
    wb.save(temp_path)
    
    return FileResponse(
        path=temp_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=output_filename,
        headers={"Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"}
    )

@router.get("/download/{session_id}/plan")
async def download_annotated_plan(session_id: str, current_user: dict = Depends(get_current_user)):
    restored = await session_manager.restore_session_if_needed(session_id, user_id=current_user["id"])
    if not restored:
        raise HTTPException(status_code=404, detail="Session not found.")
    config = {"configurable": {"thread_id": session_id}}
    state_snapshot = await costmate_graph.aget_state(config)
    if not state_snapshot.values:
        raise HTTPException(status_code=404, detail="Session not found.")
    
    annotated_path = state_snapshot.values.get("annotated_pdf_path")
    if not annotated_path:
        raise HTTPException(status_code=404, detail="Annotated plan PDF not ready or not found.")
        
    filename = f"Costmate_Annotated_Plan_{session_id}.pdf"
    media_type = "application/pdf"
    
    if annotated_path.startswith("http"):
        import httpx
        from fastapi.responses import StreamingResponse
        async def stream_external_file():
            async with httpx.AsyncClient(follow_redirects=True) as client:
                async with client.stream("GET", annotated_path) as response:
                    if response.status_code != 200:
                        raise HTTPException(status_code=404, detail="Failed to fetch annotated plan from cloud storage")
                    async for chunk in response.aiter_bytes():
                        yield chunk
        
        return StreamingResponse(
            stream_external_file(),
            media_type=media_type,
            headers={
                "Content-Disposition": f"attachment; filename={filename}",
                "Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"
            }
        )
        
    if not os.path.exists(annotated_path):
        raise HTTPException(status_code=404, detail="Annotated plan PDF not ready or not found.")
        
    return FileResponse(
        path=annotated_path, 
        media_type=media_type, 
        filename=filename,
        headers={"Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"}
    )

@router.get("/files/{session_id}/plan")
async def get_plan_image(session_id: str, current_user: dict = Depends(get_current_user)):
    restored = await session_manager.restore_session_if_needed(session_id, user_id=current_user["id"])
    if not restored:
        raise HTTPException(status_code=404, detail="Session not found.")
    config = {"configurable": {"thread_id": session_id}}
    state_snapshot = await costmate_graph.aget_state(config)
    if not state_snapshot.values:
        raise HTTPException(status_code=404, detail="Session not found.")
    page_paths = state_snapshot.values.get("uploaded_page_paths", [])
    if page_paths:
        img_path = page_paths[0]
    else:
        img_path = state_snapshot.values.get("uploaded_file_path", "")
    if not img_path:
        raise HTTPException(status_code=404, detail="Image not found")
    if img_path.startswith("http"):
        return RedirectResponse(img_path)
    if not os.path.exists(img_path):
        raise HTTPException(status_code=404, detail="Image not found")
    media_type = "application/pdf" if img_path.lower().endswith('.pdf') else "image/png"
    return FileResponse(img_path, media_type=media_type)

@router.get("/files/{session_id}/plan/pages")
async def get_plan_pages(session_id: str, current_user: dict = Depends(get_current_user)):
    restored = await session_manager.restore_session_if_needed(session_id, user_id=current_user["id"])
    if not restored:
        raise HTTPException(status_code=404, detail="Session not found.")
    config = {"configurable": {"thread_id": session_id}}
    state_snapshot = await costmate_graph.aget_state(config)
    if not state_snapshot.values:
        raise HTTPException(status_code=404, detail="Session not found.")
    page_paths = state_snapshot.values.get("uploaded_page_paths", [])
    if not page_paths:
        page_paths = [state_snapshot.values.get("uploaded_file_path", "")]
    pages = [{"page_num": i, "filename": os.path.basename(p)} for i, p in enumerate(page_paths) if p]
    return {"pages": pages, "page_count": len(pages)}
@router.get("/files/{session_id}/readme")
async def get_session_readme(session_id: str, current_user: dict = Depends(get_current_user)):
    restored = await session_manager.restore_session_if_needed(session_id, user_id=current_user["id"])
    if not restored:
        raise HTTPException(status_code=404, detail="Session not found.")
    config = {"configurable": {"thread_id": session_id}}
    state_snapshot = await costmate_graph.aget_state(config)
    if not state_snapshot.values:
        raise HTTPException(status_code=404, detail="Session not found.")
    
    qa = state_snapshot.values.get("qa_verified", {})
    floors = qa.get("floors", [])
    civil = state_snapshot.values.get("civil_quantities", {})
    columns = qa.get("columns", [])
    
    total_rooms = sum(len(f.get("rooms", [])) for f in floors) if isinstance(floors, list) else 0
    columns_count = sum(c.get("count", 1) for c in columns) if isinstance(columns, list) else 0

    return {
        "project_name": qa.get("project_name", state_snapshot.values.get("project_name", "Untitled Project")),
        "sub_work_name": qa.get("sub_work_name", "—"),
        "plan_type": qa.get("plan_type", "Architectural"),
        "original_filename": state_snapshot.values.get("original_filename", "Unknown file"),
        "has_excel": state_snapshot.values.get("excel_file_path") is not None,
        "num_floors": qa.get("num_floors", len(floors) if isinstance(floors, list) else 0),
        "total_rooms": total_rooms,
        "columns_count": columns_count,
        "floors": floors
    }

@router.get("/files/{session_id}/plan/{page_num}")
async def get_page_thumbnail(session_id: str, page_num: int, current_user: dict = Depends(get_current_user)):
    restored = await session_manager.restore_session_if_needed(session_id, user_id=current_user["id"])
    if not restored:
        raise HTTPException(status_code=404, detail="Session not found.")
    config = {"configurable": {"thread_id": session_id}}
    state_snapshot = await costmate_graph.aget_state(config)
    page_paths = state_snapshot.values.get("uploaded_page_paths", [])
    if not page_paths:
        page_paths = [state_snapshot.values.get("uploaded_file_path", "")]
    if page_num < 0 or page_num >= len(page_paths):
        raise HTTPException(status_code=404, detail="Page not found")
    img_path = page_paths[page_num]
    if not img_path:
        raise HTTPException(status_code=404, detail="Image file not found")
    if img_path.startswith("http"):
        return RedirectResponse(img_path)
    if not os.path.exists(img_path):
        raise HTTPException(status_code=404, detail="Image file not found")
    media_type = "application/pdf" if img_path.lower().endswith('.pdf') else "image/png"
    return FileResponse(img_path, media_type=media_type)

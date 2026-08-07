import openpyxl

file_path = r"c:\Users\Hp\Desktop\Costmate_v3\Assets\Costmate_Estimate_3f8a7551-a1a6-4ca0-acbb-8284c9709900.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

print("Sheet names:", wb.sheetnames)

for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n--- SHEET: {name} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    for r in range(1, min(10, ws.max_row + 1)):
        row_vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
        print(f"Row {r}: {row_vals}")
